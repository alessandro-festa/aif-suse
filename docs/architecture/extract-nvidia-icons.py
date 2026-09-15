#!/usr/bin/env python3
"""Extract the icons figure 0 needs from NVIDIA's architecture-iconography deck.

    python3 docs/architecture/extract-nvidia-icons.py \
        ~/Downloads/GTCBerlin2026_\\ ArchDiagram-Iconography.pptx

The deck is NVIDIA's and is never committed — it is an argument, not a constant.
Each icon in it is a GROUP holding one 512x512 RGBA PNG plus a caption autoshape,
so the caption is the lookup key. Captions repeat across slides (there are two
"Vault"s and two "Skill"s), so WANTED keys on (slide, caption) rather than caption
alone.

Also rasterizes the three brand lockups, which are already in this repository, via
rsvg-convert — python-pptx cannot place SVG.

Output: docs/architecture/icons/<slug>.png plus icons/MANIFEST.md.
"""
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
ICONS = HERE / "icons"

# (slide, caption) -> slug. The slide number disambiguates repeated captions.
WANTED = {
    # the human, outside the cluster. Slide 5 captions "Developer", "Admin" and
    # "Generic User" three separate groups, but all three hold the same PNG — the
    # role is in the caption, not the art. So there is one person glyph, and the
    # figures distinguish human moments with Persimmon chrome instead.
    (5, "Developer"): "developer",
    (12, "Laptop"): "laptop",
    # the agent plane, one icon per config_secops.yml _type
    (11, "Command API"): "api",
    (7, "Prompt"): "prompt",
    # NOT slide 16's "Al Agent" — that group holds a portrait photograph, an
    # in-joke rather than an icon. Deep Learning Task is the neural-net-plus-gear
    # glyph and is what the rest of NVIDIA's own diagrams use for an agent.
    (11, "Deep Learning Task"): "agent",
    (11, "Pre- and Post- Processing"): "workflow",
    (11, "Router"): "router",
    (11, "Plan"): "plan",
    (11, "Analyze Data"): "analyze-data",
    (11, "Data Integration"): "data-integration",
    (11, "NVIDIA-Accelerated Search"): "search",
    (10, "NVIDIA Skill Bundle"): "skill-bundle",
    (11, "Retriever"): "retriever",
    (7, "Enterprise Data"): "enterprise-data",
    (9, "Search Database"): "search-database",
    (9, "SQL Database"): "sql-database",
    (9, "Vector Database"): "vector-database",
    (11, "Observe"): "observe",
    (11, "Export"): "export",
    # the action plane
    (12, "Gateway"): "gateway",
    (12, "Virtual Workstation"): "virtual-workstation",
    (13, "Firewall"): "firewall",
    (13, "Vault"): "vault",
    (14, "Starfleet Service API Keys"): "starfleet-api-keys",
    (11, "Run Command"): "run-command",
    (11, "Quality Check"): "quality-check",
    (11, "Validate File"): "validate-file",
    (7, "Code"): "code",
    # models, knowledge, flywheel
    (14, "NVIDIA LLM"): "nvidia-llm",
    (14, "NVIDIA Pretrained Model"): "nvidia-pretrained-model",
    (11, "Machine Learning Task"): "machine-learning-task",
    (11, "Training Model"): "training-model",
    (14, "NVIDIA NeMo"): "nvidia-nemo",
    (14, "External Ingestion"): "ingestion",
    (13, "Model Adapter"): "model-adapter",
    (13, "Result"): "result",
    (7, "History"): "history",
    (7, "Data Lake"): "data-lake",
    # infrastructure
    (13, "Server"): "server",
    (13, "Compute"): "compute",
    (13, "Hardware"): "hardware",
    (13, "Storage"): "storage",
    (12, "On Premises"): "on-premises",
    (12, "Hybrid"): "hybrid",
}

# In-repo brand marks. Negative variants, because the slide ground is black.
# These keep their own colours: the no-green-for-SUSE rule governs diagram
# chrome, not the logos.
LOGOS = {
    "logo-suse": "ui/pkg/aif-ui/assets/SUSE_Logo-hor_L_Green-White-neg_sRGB.svg",
    "logo-nvidia": "ui/pkg/aif-ui/assets/nvidia-logo-horz-light.svg",
    "logo-ai-factory": "ui/pkg/aif-ui/assets/SUSE-AI-Factory-Logo_neg-green-horizontal.svg",
}


def caption_of(group):
    for kid in group.shapes:
        if kid.has_text_frame and kid.text_frame.text.strip():
            return kid.text_frame.text.strip().replace("\n", " ")
    return None


def picture_of(group):
    for kid in group.shapes:
        if kid.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return kid
    return None


def extract_icons(deck_path, manifest):
    prs = Presentation(deck_path)
    found = {}
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
                continue
            caption = caption_of(shape)
            slug = WANTED.get((n, caption))
            if slug is None or slug in found:
                continue
            picture = picture_of(shape)
            if picture is None:
                continue
            blob = picture.image.blob
            out = ICONS / f"{slug}.png"
            out.write_bytes(blob)
            found[slug] = True
            manifest.append(
                f"| `{slug}.png` | {len(blob):,} B | {deck_path.name} | {n} | {caption} |"
            )
            print(f"  icon  {slug:<18} slide {n:>2}  {caption}")

    missing = sorted(set(WANTED.values()) - set(found))
    if missing:
        raise SystemExit(f"missing icons, the deck changed: {missing}")


def extract_logos(manifest):
    for slug, rel in LOGOS.items():
        src = REPO / rel
        if not src.exists():
            raise SystemExit(f"missing logo source: {src}")
        out = ICONS / f"{slug}.png"
        subprocess.run(
            ["rsvg-convert", "-z", "4", "-o", str(out), str(src)], check=True
        )
        manifest.append(
            f"| `{slug}.png` | {out.stat().st_size:,} B | `{rel}` (in repo) | — | brand lockup, rsvg-convert -z 4 |"
        )
        print(f"  logo  {slug:<18} from {rel}")


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    deck = Path(sys.argv[1]).expanduser()
    if not deck.exists():
        raise SystemExit(f"no such deck: {deck}")

    ICONS.mkdir(exist_ok=True)
    manifest = []
    extract_icons(deck, manifest)
    extract_logos(manifest)

    (ICONS / "MANIFEST.md").write_text(
        "# Icon provenance\n\n"
        "Regenerate with `docs/architecture/extract-nvidia-icons.py <deck>`.\n\n"
        "The NVIDIA icons are NVIDIA brand assets, lifted from NVIDIA's own\n"
        "architecture-diagram iconography deck for GTC Berlin 2026. The source deck\n"
        "is not committed; the extractor takes its path as an argument. The three\n"
        "lockups come from this repository's own UI assets and keep their own colours.\n\n"
        "| File | Size | Source | Slide | Original caption |\n"
        "|---|---|---|---|---|\n" + "\n".join(manifest) + "\n"
    )
    print(f"\n{len(manifest)} files -> {ICONS}")


if __name__ == "__main__":
    main()
