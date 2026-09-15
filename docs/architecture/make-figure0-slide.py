#!/usr/bin/env python3
"""Build figure 0 as a single native PowerPoint slide.

    python3 docs/architecture/make-figure0-slide.py

Writes docs/architecture/figure-0-platform.pptx — one slide, 13.333 x 8.333 in
(16:10, matching make-deck.py so this can be dropped into that deck later).

Everything is a real shape, a real text run or a real icon PNG. Nothing is a
flattened diagram image, which is the whole point: the old figure 0 was an SVG
export whose agent layer was one opaque "NVIDIA AI-Q" box. Here every component
is named, and every name is a `_type` that greps out of
examples/secops-agent-factory/configs/config_secops.yml.

Run extract-nvidia-icons.py first; this reads icons/ and will not invent art.

COLOUR RULE, enforced by the constants below and by verify-figure0-slide.py:
green means NVIDIA and nothing else. SUSE is Waterhole blue. The logos are
pictures and keep their own colours — the rule governs diagram chrome.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ICONS = HERE / "icons"
OUT = HERE / "figure-0-platform.pptx"

# --- palette -----------------------------------------------------------------
# NVIDIA: from the GTC Berlin 2026 theme (lt2 = 76B900, dk2 = 000000).
# SUSE:   from brand.suse.com/design-language#color, Waterhole and the neutral
#         ramp. Jungle #30BA78 is deliberately absent — it is the colour that
#         made SUSE and NVIDIA indistinguishable in the old figure.
NVIDIA_GREEN = RGBColor(0x76, 0xB9, 0x00)   # NVIDIA, and only NVIDIA
HINGE_LIME = RGBColor(0xC8, 0xE6, 0x00)     # the hinge — both halves are NVIDIA
WATERHOLE = RGBColor(0x24, 0x53, 0xFF)      # SUSE — supplied and packaged
WATERHOLE_MID = RGBColor(0x3C, 0x8E, 0xEF)  # SUSE strokes and small type
WATERHOLE_LIGHT = RGBColor(0x81, 0xAE, 0xFC)  # SUSE — platform prerequisite
MIDNIGHT_DEEP = RGBColor(0x0A, 0x11, 0x2B)  # SUSE panel fill
PERSIMMON = RGBColor(0xFE, 0x7C, 0x3F)      # the human path
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FOG = RGBColor(0xEF, 0xEF, 0xEF)
GREY_BA = RGBColor(0xBA, 0xBA, 0xBA)
GREY_99 = RGBColor(0x99, 0x99, 0x99)
GREY_6F = RGBColor(0x6F, 0x6F, 0x6F)
GREY_3E = RGBColor(0x3E, 0x3E, 0x3E)
PANEL = RGBColor(0x0B, 0x0B, 0x0B)          # NVIDIA panel fill, neutral

# NVIDIA Sans is the deck's theme font but is installed neither here nor in
# Google Slides, so it would fall back to something arbitrary. Arial is the
# closest safe match everywhere, the same call make-deck.py makes.
FONT = "Arial"

W, H = Inches(13.333), Inches(8.333)


# --- primitives --------------------------------------------------------------
def box(slide, x, y, w, h, fill=None, line=None, lw=Pt(1.0), dash=False, r=0.04):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = r
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = lw
        if dash:
            s.line.dash_style = 4  # msoLineDash
    s.shadow.inherit = False       # PowerPoint's default shadow looks wrong on black
    s.text_frame.text = ""
    return s


def label(slide, x, y, w, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, pt, bold, colour) — one paragraph each, tight.

    No height: a textbox on a black ground has no fill or border, so its box is
    invisible and only the text position matters. Fixing it here keeps every
    call site down to the two numbers that do matter.
    """
    tb = slide.shapes.add_textbox(x, y, w, Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (text, size, bold, colour) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        p.line_spacing = 0.94
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = colour
            run.font.name = FONT
    return tb


def icon(slide, name, x, y, size):
    return slide.shapes.add_picture(str(ICONS / f"{name}.png"), x, y, height=size)


def logo(slide, name, x, y, height):
    return slide.shapes.add_picture(str(ICONS / f"{name}.png"), x, y, height=height)


def arrow(slide, x1, y1, x2, y2, colour, lw=Pt(1.25), dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = colour
    c.line.width = lw
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    end = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(end)
    c.shadow.inherit = False
    return c


def item(slide, x, y, w, ico, name, sub, name_colour, sub_colour, ih=Inches(0.24)):
    """One icon + two-line entry. The unit the whole figure is built from."""
    icon(slide, ico, x, y + Emu(int(ih * 0.06)), ih)
    tx = x + ih + Inches(0.07)
    label(slide, tx, y, w - (tx - x), [
        (name, 7.5, True, name_colour),
        (sub, 5.8, False, sub_colour),
    ])


def caps(slide, x, y, w, text, colour, size=7.0):
    return label(slide, x, y, w, [(text, size, True, colour)])


# --- content -----------------------------------------------------------------
# Every `name` below is a literal key from config_secops.yml. Nothing invented:
# `_type` is a plugin-registry lookup, so a name that is not upstream does not
# start. See aiq-component-inventory.md for the line numbers.
AGENT_PLANE = [
    ("api", "aiq_api", "front_end · REST + async Job API · Postgres job store", NVIDIA_GREEN),
    ("workflow", "chat_deepresearcher_agent", "the workflow graph · escalation · clarifier · checkpointed", NVIDIA_GREEN),
    ("router", "intent_classifier", "meta vs research, shallow vs deep · Lightning", NVIDIA_GREEN),
    ("plan", "clarifier_agent", "THE HUMAN PLAN GATE · max_turns 3 · Ultra", PERSIMMON),
    ("analyze-data", "shallow_research_agent", "triage · bounded llm turns and tool iterations", NVIDIA_GREEN),
    ("agent", "deep_research_agent", "orchestrator · source_router · planner · researcher ×N · writer", NVIDIA_GREEN),
    ("skill-bundle", "deep_research_skills", "role→skill grants · require_sandbox: [research]", HINGE_LIME),
    ("virtual-workstation", "deep_research_sandbox", "provider: openshell · network: allowlist", HINGE_LIME),
    ("data-integration", "data_source_registry", "one source: the vulnerability corpus", NVIDIA_GREEN),
    ("retriever", "knowledge_retrieval", "backend: foundational_rag · top_k 10 — the retrieval seam", NVIDIA_GREEN),
    ("sql-database", "checkpoint_db", "Postgres · a job outlives the connection that made it", NVIDIA_GREEN),
    ("export", "telemetry · console", "OTLP out to SUSE Observability :4318", NVIDIA_GREEN),
]

OPENSHELL = [
    ("gateway", "Gateway", "eBPF-backed policy, enforced out of process"),
    ("firewall", "L7 egress allow-list", "FAIL-CLOSED · deny unless named"),
    ("vault", "Credential injection", "at the boundary · full session recording"),
]

SANDBOXES = [
    ("run-command", "Remediation engineer", "Ultra · the only agent that writes",
     "git clone / push · buildah build · zypper patch · pytest"),
    ("quality-check", "Validation agent", "Lightning · grades the fix",
     "rebuild image · SUSE Security scan · scan diff · gate the PR"),
]

MODEL_PLANE = [
    ("nvidia-llm", "Nemotron 3 Ultra", "550B-A55B · plans, routes, writes, verifies"),
    ("nvidia-pretrained-model", "Nemotron 3.5 Lightning", "30B-A3B · executes · most of the tokens"),
    ("machine-learning-task", "Nemotron 3 Embed 1B + Rerank VL 1B", "retrieval · 2.2.2"),
    ("nvidia-nemo", "NeMo Guardrails 25.6.0", "intercepts every call above"),
]

KNOWLEDGE_PLANE = [
    ("search", "rag-server", "hybrid search, sparse-weighted — a NEVRA must match literally"),
    ("ingestion", "ingestor-server · nv-ingest", "table structure + page elements on · OCR off"),
    ("vector-database", "Elasticsearch", "the vector store"),
    ("enterprise-data", "the corpus", "CVE + SUSE-SU feeds · SBOMs · SUSE Security scan exports"),
]

FLYWHEEL = [
    ("history", "Relay", "ATIF traces · ships in the sandbox image"),
    ("data-lake", "Data Store", "what it actually did"),
    ("result", "Evaluator", "LLM-as-judge on accepted trajectories"),
    ("model-adapter", "Customizer", "LoRA on the 30B, folded back weekly"),
]

INFRA = [
    ("server", "SUSE Linux Enterprise", "SLES 16 · SL Micro on every node", WATERHOLE_MID),
    ("hybrid", "Rancher Prime 2.14.3", "lifecycle · ClusterRepos · Fleet GitOps", WATERHOLE_MID),
    ("on-premises", "RKE2 1.35.6", "CIS-hardened · FIPS-capable · air-gap installable", WATERHOLE_MID),
    ("storage", "SUSE Storage 1.11.3", "model cache · vector store · checkpoints", WATERHOLE_MID),
    ("compute", "NVIDIA GPU Operator v26.3.1", "driver 595 from registry.suse.com · CDI by NRI", NVIDIA_GREEN),
    ("hardware", "NVIDIA accelerated compute", "MIG and DRA · Network Operator · Run:ai optional", NVIDIA_GREEN),
]

LEGEND = [
    (NVIDIA_GREEN, "NVIDIA — models, agent toolkit, sandbox runtime, GPU stack"),
    (WATERHOLE, "SUSE — OS, Kubernetes, packaging, GitOps"),
    (WATERHOLE_LIGHT, "SUSE platform prerequisite — integrated, not installed"),
    (HINGE_LIME, "the hinge: agents in OpenShell sandboxes"),
    (PERSIMMON, "human"),
]


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BLACK

    L, R = Inches(0.30), Inches(13.03)
    CW = R - L

    # --- title band ----------------------------------------------------------
    logo(slide, "logo-suse", L, Inches(0.17), Inches(0.26))
    rule = box(slide, Inches(1.62), Inches(0.155), Inches(0.012), Inches(0.30), fill=GREY_3E)
    rule.line.fill.background()
    logo(slide, "logo-nvidia", Inches(1.78), Inches(0.20), Inches(0.20))
    label(slide, Inches(3.40), Inches(0.13), Inches(9.63), [
        ("Sovereign Agentic SecOps — every component, named", 15, True, WHITE),
        ("Autonomous vulnerability remediation. Long-running, specialized agents that are allowed to act.", 7.5, False, GREY_99),
    ])

    # --- analyst workstation -------------------------------------------------
    wy, wh = Inches(0.66), Inches(0.66)
    box(slide, L, wy, CW, wh, fill=BLACK, line=PERSIMMON, lw=Pt(1.2))
    icon(slide, "developer", L + Inches(0.10), wy + Inches(0.14), Inches(0.38))
    label(slide, L + Inches(0.55), wy + Inches(0.13), Inches(2.95), [
        ("ANALYST WORKSTATION — OUTSIDE THE CLUSTER", 7.5, True, PERSIMMON),
        ("No kubeconfig. No cluster credentials. No kubectl.", 6.0, False, GREY_BA),
        ("Each of the three is a human gate.", 6.0, False, GREY_99),
    ])
    ways = [
        ("laptop", "nemoclaw connect", "SSH over HTTP CONNECT to the gateway — attach to a running sandbox and watch"),
        ("prompt", "aiq-frontend web UI", "submit a finding, accept or narrow the plan, read the cited brief"),
        ("code", "git review", "the only path to production. The agent opens the PR; it cannot merge it."),
    ]
    for i, (ico, name, sub) in enumerate(ways):
        x = Inches(3.78) + i * Inches(3.12)
        icon(slide, ico, x, wy + Inches(0.15), Inches(0.32))
        label(slide, x + Inches(0.35), wy + Inches(0.14), Inches(2.62), [
            (name, 7.5, True, PERSIMMON),
            (sub, 5.8, False, GREY_99),
        ])

    # --- RKE2 cluster frame --------------------------------------------------
    cy, cbot = Inches(1.44), Inches(7.20)
    box(slide, L, cy, CW, cbot - cy, fill=MIDNIGHT_DEEP, line=WATERHOLE, lw=Pt(1.4))
    label(slide, L + Inches(0.12), cy + Inches(0.07), Inches(7.0), [
        ("RKE2 CLUSTER  ·  managed by Rancher Prime  ·  customer estate, sovereign, air-gappable", 8, True, WATERHOLE_LIGHT),
    ])
    label(slide, R - Inches(4.6), cy + Inches(0.07), Inches(4.48), [
        ("ONE SUPPORT CONTRACT ACROSS EVERY COMPONENT BELOW", 6.5, True, WATERHOLE_MID),
    ], align=PP_ALIGN.RIGHT)
    arrow(slide, Inches(5.2), wy + wh, Inches(5.2), cy, PERSIMMON)
    label(slide, Inches(5.30), wy + wh + Inches(0.005), Inches(2.6), [
        ("submit · approve the plan · review the PR", 5.8, False, PERSIMMON)])

    # --- AI Factory frame ----------------------------------------------------
    fx, fy = L + Inches(0.14), Inches(1.72)
    fw, fbot = CW - Inches(0.28), Inches(6.60)
    box(slide, fx, fy, fw, fbot - fy, fill=BLACK, line=WATERHOLE_MID, lw=Pt(1.2))
    logo(slide, "logo-ai-factory", fx + Inches(0.12), fy + Inches(0.08), Inches(0.15))
    label(slide, fx + Inches(1.68), fy + Inches(0.055), Inches(6.0), [
        ("with NVIDIA  —  Blueprints → AIWorkloads → Fleet HelmOps · three Blueprints, three catalog apps", 6.8, True, GREY_99),
    ])

    ix = fx + Inches(0.14)
    iw = fw - Inches(0.28)

    # --- agent plane ---------------------------------------------------------
    ay, abot = fy + Inches(0.32), Inches(4.42)
    aw = Inches(7.15)
    box(slide, ix, ay, aw, abot - ay, fill=PANEL, line=NVIDIA_GREEN, lw=Pt(1.2))
    label(slide, ix + Inches(0.12), ay + Inches(0.08), Inches(5.4), [
        ("AGENT PLANE — NVIDIA NeMo Agent Toolkit", 9, True, NVIDIA_GREEN),
        ("packaged today as AI-Q 2.2 · chart nvidia-blueprints/aiq2-web 2.2.1 · ns-secops-agents", 6.0, False, GREY_6F),
    ])
    for i, (ico, name, sub, colour) in enumerate(AGENT_PLANE):
        col, row = divmod(i, 6)
        x = ix + Inches(0.12) + col * Inches(3.50)
        y = ay + Inches(0.44) + row * Inches(0.325)
        item(slide, x, y, Inches(3.42), ico, name, sub, colour, GREY_99)

    # The one seam that moves if NeMo Retriever absorbs AI-Q's retrieval half.
    # Drawn between the two panels rather than across the rows, so it reads as a
    # dependency and not as a strike-through.
    seam_x = ix + Inches(4.30)
    arrow(slide, seam_x, abot, seam_x, Inches(4.54), NVIDIA_GREEN, lw=Pt(1.5))
    label(slide, seam_x + Inches(0.07), abot - Inches(0.01), Inches(3.2),
          [("foundational_rag over HTTP — the only retrieval seam", 5.8, True, NVIDIA_GREEN)])

    # --- OpenShell -----------------------------------------------------------
    ox = ix + aw + Inches(0.12)
    ow = Inches(1.92)
    box(slide, ox, ay, ow, abot - ay, fill=PANEL, line=HINGE_LIME, lw=Pt(1.2), dash=True)
    label(slide, ox + Inches(0.10), ay + Inches(0.08), Inches(1.72), [
        ("NVIDIA OpenShell", 8.5, True, HINGE_LIME),
        ("the enforcement boundary", 6.0, False, GREY_6F),
    ])
    for i, (ico, name, sub) in enumerate(OPENSHELL):
        item(slide, ox + Inches(0.10), ay + Inches(0.44) + i * Inches(0.46), Inches(1.74),
             ico, name, sub, HINGE_LIME, GREY_99)
    label(slide, ox + Inches(0.10), ay + Inches(1.90), Inches(1.74), [
        ("No kubeconfig is ever mounted. The agent's blast radius is one policy file, "
         "and Fleet reconciles that file.", 5.8, False, GREY_6F),
    ])

    # --- sandboxes -----------------------------------------------------------
    sx = ox + ow + Inches(0.12)
    sw = ix + iw - sx
    box(slide, sx, ay, sw, abot - ay, fill=PANEL, line=HINGE_LIME, lw=Pt(1.2))
    label(slide, sx + Inches(0.10), ay + Inches(0.08), sw - Inches(0.20), [
        ("SECURE AGENT SANDBOXES", 8.5, True, HINGE_LIME),
        ("SUSE SLE BCI 16 · no public PyPI · one writable git remote", 6.0, False, GREY_6F),
    ])
    for i, (ico, name, sub, skills) in enumerate(SANDBOXES):
        y = ay + Inches(0.50) + i * Inches(0.86)
        item(slide, sx + Inches(0.10), y, sw - Inches(0.20), ico, name, sub, WHITE, GREY_99)
        label(slide, sx + Inches(0.38), y + Inches(0.30), sw - Inches(0.48), [
            ("SKILLS  " + skills, 5.8, False, HINGE_LIME)])
    label(slide, sx + Inches(0.10), abot - Inches(0.32), sw - Inches(0.20), [
        ("One sandbox per agent per job · scaled 0→1 on connect-intent, torn down after · "
         "snapshot and restore, not restart.", 5.8, False, GREY_6F),
    ])
    # the hinge, drawn at the deep_research_sandbox row: declaration → boundary → sandbox
    hy = ay + Inches(0.87)
    arrow(slide, ix + aw, hy, ox, hy, HINGE_LIME)
    arrow(slide, ox + ow, hy, sx, hy, HINGE_LIME)

    # --- model / knowledge / flywheel ---------------------------------------
    py, pbot = Inches(4.54), Inches(6.46)
    pw = (iw - Inches(0.24)) / 3
    planes = [
        ("MODEL PLANE", "NVIDIA NIM Operator 3.1.2 · NIMCache / NIMService · LiteLLM virtual keys",
         MODEL_PLANE, "Ultra plans, Lightning executes — and that split is the reason the flywheel pays."),
        ("KNOWLEDGE PLANE", "NVIDIA NeMo Retriever · NVIDIA RAG Blueprint v2.6.0 · ns-secops-knowledge",
         KNOWLEDGE_PLANE, "The only retrieval in the architecture. If Retriever absorbs AI-Q's retrieval half, this band is what moves."),
        ("DATA FLYWHEEL", "NVIDIA NeMo · accepted trajectories and human verdicts, graded and folded back",
         FLYWHEEL, "NeMo Microservices is superseded by NeMo Platform from 1 October 2026."),
    ]
    for i, (title, sub, rows, foot) in enumerate(planes):
        x = ix + i * (pw + Inches(0.12))
        box(slide, x, py, pw, pbot - py, fill=PANEL, line=NVIDIA_GREEN, lw=Pt(1.0))
        label(slide, x + Inches(0.10), py + Inches(0.07), pw - Inches(0.20), [
            (title, 8.5, True, NVIDIA_GREEN),
            (sub, 5.8, False, GREY_6F),
        ])
        for j, (ico, name, rsub) in enumerate(rows):
            item(slide, x + Inches(0.10), py + Inches(0.40) + j * Inches(0.325),
                 pw - Inches(0.20), ico, name, rsub, WHITE, GREY_99)
        label(slide, x + Inches(0.10), pbot - Inches(0.22), pw - Inches(0.20), [
            (foot, 5.8, False, GREY_6F)])

    # --- SUSE prerequisites --------------------------------------------------
    qy, qh = Inches(6.68), Inches(0.48)
    qw = (CW - Inches(0.28) - Inches(0.12)) / 2
    prereqs = [
        ("validate-file", "SUSE Security (NeuVector)",
         "Registry scanning · runtime enforcer · admission control · REST API :10443, read-only to the agent",
         "Raises the signal and grades the fix. The agent may read the scanner; policy denies it the scanner's own rules."),
        ("observe", "SUSE Observability",
         "OTLP :4318 from the agents and the RAG server · :4317 gRPC from the gateway · MCP topology · AI assistant OFF",
         "One topology over models, retrieval, agents, sandboxes and the cluster underneath all of them."),
    ]
    for i, (ico, name, sub, foot) in enumerate(prereqs):
        x = fx + i * (qw + Inches(0.12))
        box(slide, x, qy, qw, qh, fill=MIDNIGHT_DEEP, line=WATERHOLE_LIGHT, lw=Pt(1.0), dash=True)
        icon(slide, ico, x + Inches(0.09), qy + Inches(0.12), Inches(0.24))
        label(slide, x + Inches(0.40), qy + Inches(0.05), qw - Inches(0.50), [
            (name + "   —   CUSTOMER PREREQUISITE, INTEGRATED NOT INSTALLED", 7.5, True, WATERHOLE_LIGHT),
            (sub, 5.8, False, GREY_99),
            (foot, 5.8, False, GREY_6F),
        ])

    # --- infrastructure ------------------------------------------------------
    ny, nh = Inches(7.40), Inches(0.46)
    caps(slide, L, ny - Inches(0.155), Inches(11.0),
         "INFRASTRUCTURE AND ACCELERATED COMPUTE — CO-ENGINEERED AND SHIPPED AS ONE RELEASE MANIFEST", GREY_6F, 6.5)
    nw = (CW - 5 * Inches(0.10)) / 6
    for i, (ico, name, sub, colour) in enumerate(INFRA):
        x = L + i * (nw + Inches(0.10))
        box(slide, x, ny, nw, nh, fill=MIDNIGHT_DEEP if colour is WATERHOLE_MID else PANEL,
            line=colour, lw=Pt(1.0))
        icon(slide, ico, x + Inches(0.07), ny + Inches(0.13), Inches(0.24))
        label(slide, x + Inches(0.36), ny + Inches(0.08), nw - Inches(0.44), [
            (name, 7, True, colour),
            (sub, 5.5, False, GREY_99),
        ])

    # --- legend and status ---------------------------------------------------
    ly = Inches(7.96)
    x = L
    for colour, text in LEGEND:
        sw_ = box(slide, x, ly + Inches(0.022), Inches(0.10), Inches(0.10), fill=colour)
        sw_.line.fill.background()
        tb = label(slide, x + Inches(0.15), ly, Inches(3.0), [(text, 6.0, False, GREY_99)])
        x += Inches(0.15) + Inches(0.045) * len(text) + Inches(0.14)
    label(slide, L, ly + Inches(0.20), CW, [
        ("Reference architecture. OpenShell, NemoClaw and the sandbox lifecycle are verified on a live SUSE cluster; "
         "the SecOps agent layer, the data flywheel and the three portfolio integrations are designed, not built.",
         5.8, False, GREY_6F),
    ])

    prs.save(OUT)
    print(f"{len(slide.shapes)} shapes -> {OUT}")


if __name__ == "__main__":
    build()
