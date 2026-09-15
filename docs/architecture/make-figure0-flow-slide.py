#!/usr/bin/env python3
"""Build the simplified figure 0 — the data flow, not the inventory.

    python3 docs/architecture/make-figure0-flow-slide.py

Writes docs/architecture/figure-0-flow.pptx — one slide, 13.333 x 8.333 in,
the same canvas as make-figure0-slide.py so both drop into the same deck.

figure-0-platform.pptx names every component; that is its job and it keeps it.
This slide answers the other question: what happens to a query? It is the same
architecture at a quarter of the text, in stacked layers — a control band across
the top, the agent plane and the sandbox boundary side by side in the middle,
the AI Factory underneath, and a dashed return that closes the flywheel.

The unit here is tile(): one large icon with two short lines centred under it.
If a tile needs a third line, the slide is drifting back into an inventory.

Palette and primitives are imported from make-figure0-slide.py rather than
copied, so there is one definition of the colour rule — green means NVIDIA —
and verify-figure0-slide.py lints both slides against it.
"""
import importlib.util
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "figure-0-flow.pptx"

sys.dont_write_bytecode = True
_spec = importlib.util.spec_from_file_location("figure0", HERE / "make-figure0-slide.py")
base = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(base)

box, label, icon, logo, caps = base.box, base.label, base.icon, base.logo, base.caps
NVIDIA_GREEN, HINGE_LIME = base.NVIDIA_GREEN, base.HINGE_LIME
WATERHOLE, WATERHOLE_MID, WATERHOLE_LIGHT = base.WATERHOLE, base.WATERHOLE_MID, base.WATERHOLE_LIGHT
PERSIMMON, PANEL = base.PERSIMMON, base.PANEL
BLACK, WHITE = base.BLACK, base.WHITE
GREY_BA, GREY_99, GREY_6F, GREY_3E = base.GREY_BA, base.GREY_99, base.GREY_6F, base.GREY_3E
W, H = base.W, base.H


# --- primitives this slide adds ----------------------------------------------
def seg(slide, x1, y1, x2, y2, colour, lw=Pt(1.1), dash=False, head=True):
    """arrow() with the head made optional, because the feedback return is an
    elbow and MSO_CONNECTOR.STRAIGHT cannot bend — it is drawn as three
    segments and only the last one carries a head.

    The int() calls are not cosmetic. Column widths here are Emu divisions and
    come out as floats; add_shape and add_textbox coerce, but add_connector
    writes its arguments straight through, and a single `x="11804904.0"` is an
    invalid ST_Coordinate that makes PowerPoint refuse to open the file.
    """
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = colour
    c.line.width = lw
    if dash:
        c.line.dash_style = 4
    if head:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))
    c.shadow.inherit = False
    return c


def tile(slide, cx, y, w, ico, name, sub, name_colour, sub_colour,
         ih=Inches(0.46), name_pt=7.5, sub_pt=6.0):
    """One large icon centred on cx, two centred lines under it. The unit.

    icon() sizes by height only so aspect varies, so the picture is placed and
    then shifted by its own measured width — there is no width to predict.
    """
    pic = icon(slide, ico, cx, y, ih)
    pic.left = int(cx - pic.width / 2)
    label(slide, int(cx - w / 2), int(y + ih + Inches(0.07)), w, [
        (name, name_pt, True, name_colour),
        (sub, sub_pt, False, sub_colour),
    ], align=PP_ALIGN.CENTER)
    return pic


# --- content -----------------------------------------------------------------
# The loop that surrounds every job. Two of the five are customer prerequisites
# SUSE integrates rather than installs; one of the five is a human.
CONTROL = [
    ("DETECT", "validate-file", "SUSE Security", "registry scan · runtime enforcer", WATERHOLE_LIGHT),
    ("GOVERN", "nvidia-nemo", "NeMo Guardrails", "intercepts every model call", NVIDIA_GREEN),
    ("APPROVE", "developer", "The human", "the plan gate, then the PR", PERSIMMON),
    ("OBSERVE", "observe", "SUSE Observability", "OTLP :4318 · one topology", WATERHOLE_LIGHT),
    ("LEARN", "training-model", "Data flywheel", "the verdict becomes a label", NVIDIA_GREEN),
]

# The job, in order. Four steps, and the second one stops for a person.
AGENT_STEPS = [
    ("router", "Classify and triage", "intent_classifier · one job, not fourteen", NVIDIA_GREEN),
    ("plan", "The plan gate", "clarifier_agent · holds for a human", PERSIMMON),
    ("agent", "Research fan-out", "deep_research_agent · workers ×N", NVIDIA_GREEN),
    ("quality-check", "The cited brief", "citation verifier · fails closed", NVIDIA_GREEN),
]

AGENT_PLUMBING = [
    ("api", "aiq_api", "REST + async job API · a job outlives its connection"),
    ("retriever", "knowledge_retrieval", "foundational_rag over HTTP — the only retrieval seam"),
]

# The two agents allowed to change the world, and the boundary they work behind.
WRITERS = [
    ("run-command", "Remediation engineer", "patch · build · pytest · opens the PR"),
    ("analyze-data", "Validation agent", "rebuild · rescan · diffs it on the PR"),
]

BOUNDARY = [
    ("gateway", "Gateway", "policy enforced out of process"),
    ("firewall", "L7 egress allow-list", "FAIL-CLOSED · deny unless named"),
    ("starfleet-api-keys", "Credential injection", "the token never enters the sandbox"),
]

EXIT = [
    ("code", "Pull request", "the agent opens it · it cannot merge, approve or unprotect"),
    ("data-integration", "Fleet applies", "GitOps · never kubectl"),
]

PLANES = [
    ("MODELS", "NVIDIA NIM Operator · NeMo Guardrails above", [
        ("nvidia-llm", "Nemotron 3 Ultra", "plans · writes · verifies"),
        ("nvidia-pretrained-model", "Nemotron 3.5 Lightning", "executes · most of the tokens"),
        ("machine-learning-task", "Embed + Rerank", "retrieval models"),
    ]),
    ("KNOWLEDGE", "NVIDIA NeMo Retriever · RAG Blueprint v2.6.0", [
        ("search", "rag-server", "hybrid, sparse-weighted"),
        ("ingestion", "nv-ingest", "tables + page elements"),
        ("vector-database", "Elasticsearch", "the vector store"),
        ("enterprise-data", "the corpus", "CVE · SUSE-SU · SBOMs"),
    ]),
    ("DATA FLYWHEEL", "NVIDIA NeMo · graded and folded back weekly", [
        ("history", "Relay", "ATIF traces"),
        ("data-lake", "Data Store", "what it actually did"),
        ("result", "Evaluator", "LLM-as-judge"),
        ("model-adapter", "Customizer", "LoRA on the 30B"),
    ]),
]

INFRA = [
    ("server", "SUSE Linux Enterprise", "SLES 16 · SL Micro", WATERHOLE_MID),
    ("hybrid", "Rancher Prime", "lifecycle · Fleet GitOps", WATERHOLE_MID),
    ("on-premises", "RKE2", "CIS-hardened · air-gappable", WATERHOLE_MID),
    ("storage", "SUSE Storage", "model cache · vector store", WATERHOLE_MID),
    ("compute", "NVIDIA GPU Operator", "driver from registry.suse.com", NVIDIA_GREEN),
]

LEGEND = [
    (NVIDIA_GREEN, "NVIDIA"),
    (WATERHOLE, "SUSE"),
    (WATERHOLE_LIGHT, "SUSE prerequisite"),
    (HINGE_LIME, "the hinge"),
    (PERSIMMON, "human"),
]


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BLACK

    L, R = Inches(0.30), Inches(13.03)
    CW = R - L

    # --- title band ----------------------------------------------------------
    logo(slide, "logo-suse", L, Inches(0.17), Inches(0.26))
    rule = box(slide, Inches(1.62), Inches(0.155), Inches(0.012), Inches(0.30), fill=GREY_3E)
    rule.line.fill.background()
    logo(slide, "logo-nvidia", Inches(1.78), Inches(0.20), Inches(0.20))
    label(slide, Inches(3.40), Inches(0.13), Inches(9.63), [
        ("Sovereign Agentic SecOps — a finding becomes a reviewed pull request", 15, True, WHITE),
        ("One query. Two human gates. The agent proposes a change to a manifest in git; Fleet applies it.",
         7.5, False, GREY_99),
    ])

    # --- control band --------------------------------------------------------
    by, bbot = Inches(0.62), Inches(1.76)
    box(slide, L, by, CW, bbot - by, fill=PANEL, line=WATERHOLE_LIGHT, lw=Pt(1.1), dash=True)
    caps(slide, L + Inches(0.12), by + Inches(0.045), Inches(8.0),
         "THE LOOP AROUND EVERY JOB — two SUSE prerequisites, one guardrail, one human, one flywheel",
         GREY_99, size=7.0)
    cw = (CW - 4 * Inches(0.10)) / 5
    for i, (head, ico, name, sub, colour) in enumerate(CONTROL):
        x = L + i * (cw + Inches(0.10))
        cx = int(x + cw / 2)
        label(slide, x, by + Inches(0.255), int(cw), [(head, 6.8, True, colour)],
              align=PP_ALIGN.CENTER)
        tile(slide, cx, by + Inches(0.38), int(cw), ico, name, sub, colour, GREY_99, ih=Inches(0.30))

    # --- the two arrows between the layers -----------------------------------
    my, mbot = Inches(2.06), Inches(5.30)
    seg(slide, Inches(3.30), bbot, Inches(3.30), my, WATERHOLE_MID)
    label(slide, Inches(3.40), bbot + Inches(0.02), Inches(3.2),
          [("the finding · the scan export · policy", 5.8, False, WATERHOLE_MID)])
    seg(slide, Inches(9.60), my, Inches(9.60), bbot, WATERHOLE_MID)
    label(slide, Inches(9.70), bbot + Inches(0.02), Inches(3.2),
          [("telemetry · OTLP :4318 · audit", 5.8, False, WATERHOLE_MID)])

    # --- the human, outside the cluster --------------------------------------
    tile(slide, Inches(0.76), Inches(2.62), Inches(0.92), "laptop", "Analyst",
         "no kubeconfig", PERSIMMON, GREY_99, ih=Inches(0.44), sub_pt=5.8)
    seg(slide, Inches(1.24), Inches(2.84), Inches(1.40), Inches(2.84), PERSIMMON)

    # --- agent plane ---------------------------------------------------------
    ax, aw = Inches(1.40), Inches(5.44)
    box(slide, ax, my, aw, mbot - my, fill=PANEL, line=NVIDIA_GREEN, lw=Pt(1.2))
    label(slide, ax + Inches(0.12), my + Inches(0.06), Inches(5.2), [
        ("AGENT PLANE — NVIDIA NeMo Agent Toolkit", 9, True, NVIDIA_GREEN),
        ("AI-Q 2.2 · ns-secops-agents · eight agents read, two write", 6.0, False, GREY_6F),
    ])
    ix, iw = ax + Inches(0.12), aw - Inches(0.24)

    box(slide, ix, Inches(2.46), iw, Inches(1.24), line=GREY_3E, lw=Pt(0.75))
    caps(slide, ix + Inches(0.08), Inches(2.52), Inches(4.0), "THE JOB, IN ORDER", GREY_6F, size=6.5)
    sw = (iw - Inches(0.16)) / 4
    for i, (ico, name, sub, colour) in enumerate(AGENT_STEPS):
        cx = int(ix + Inches(0.08) + i * sw + sw / 2)
        tile(slide, cx, Inches(2.72), int(sw), ico, name, sub, colour, GREY_99)
        if i:
            seg(slide, int(cx - sw / 2 - Inches(0.06)), Inches(2.95),
                int(cx - sw / 2 + Inches(0.02)), Inches(2.95),
                PERSIMMON if colour is PERSIMMON else NVIDIA_GREEN, lw=Pt(0.9))

    box(slide, ix, Inches(3.82), iw, Inches(0.96), line=GREY_3E, lw=Pt(0.75), dash=True)
    caps(slide, ix + Inches(0.08), Inches(3.88), Inches(4.0), "THE PLUMBING UNDERNEATH", GREY_6F, size=6.5)
    pw = iw / 2
    for i, (ico, name, sub) in enumerate(AGENT_PLUMBING):
        tile(slide, int(ix + i * pw + pw / 2), Inches(4.06), int(pw), ico, name, sub,
             NVIDIA_GREEN, GREY_99, ih=Inches(0.30), sub_pt=5.8)

    label(slide, ix, Inches(4.92), iw, [
        ("Every step above is a `_type` in config_secops.yml. Nothing here can reach the cluster API.",
         5.8, False, GREY_6F)])

    # --- the boundary and the sandboxes --------------------------------------
    ox, ow = Inches(7.46), Inches(4.16)
    box(slide, ox, my, ow, mbot - my, fill=PANEL, line=HINGE_LIME, lw=Pt(1.2))
    icon(slide, "virtual-workstation", ox + Inches(0.12), my + Inches(0.04), Inches(0.20))
    label(slide, ox + Inches(0.38), my + Inches(0.045), Inches(3.66), [
        ("SECURE AGENT SANDBOXES", 9, True, HINGE_LIME),
        ("SUSE SLE BCI 16 · no public PyPI · one writable git remote", 6.0, False, GREY_6F),
    ])
    jx, jw = ox + Inches(0.12), ow - Inches(0.24)

    tw = jw / 2
    for i, (ico, name, sub) in enumerate(WRITERS):
        tile(slide, int(jx + i * tw + tw / 2), Inches(2.62), int(tw), ico, name, sub,
             HINGE_LIME, GREY_99)

    box(slide, jx, Inches(3.62), jw, Inches(1.16), line=HINGE_LIME, lw=Pt(0.9), dash=True)
    caps(slide, jx + Inches(0.08), Inches(3.68), Inches(3.6),
         "NVIDIA OPENSHELL — ENFORCED OUT OF PROCESS", HINGE_LIME, size=6.5)
    bw = jw / 3
    for i, (ico, name, sub) in enumerate(BOUNDARY):
        tile(slide, int(jx + i * bw + bw / 2), Inches(3.88), int(bw), ico, name, sub,
             HINGE_LIME, GREY_99, ih=Inches(0.34), name_pt=7.0, sub_pt=5.8)

    label(slide, jx, Inches(4.92), jw, [
        ("A jailbroken agent cannot widen this: the supervisor and the egress proxy are separate processes.",
         5.8, False, GREY_6F)])

    # left plane hands off to the boundary
    seg(slide, ax + aw, Inches(2.84), ox, Inches(2.84), HINGE_LIME)
    label(slide, ax + aw - Inches(0.02), Inches(2.52), Inches(0.66), [
        ("plan approved", 5.5, False, HINGE_LIME),
        ("skills granted", 5.5, False, HINGE_LIME),
    ], align=PP_ALIGN.CENTER)

    # --- the way out ---------------------------------------------------------
    ex, ew = Inches(11.84), Inches(1.19)
    ecx = int(ex + ew / 2)
    label(slide, ex, my + Inches(0.06), ew, [("TO PRODUCTION", 6.8, True, PERSIMMON)],
          align=PP_ALIGN.CENTER)
    seg(slide, ox + ow, Inches(2.84), ex, Inches(2.84), PERSIMMON)
    for i, (ico, name, sub) in enumerate(EXIT):
        tile(slide, ecx, Inches(2.50) + i * Inches(1.10), ew, ico, name, sub,
             PERSIMMON, GREY_99, ih=Inches(0.42), sub_pt=5.5)
    seg(slide, ecx, Inches(3.44), ecx, Inches(3.58), PERSIMMON, lw=Pt(0.9))

    # --- the AI Factory underneath -------------------------------------------
    py, pbot = Inches(5.44), Inches(7.46)
    box(slide, L, py, CW, pbot - py, fill=BLACK, line=WATERHOLE_MID, lw=Pt(1.2))
    logo(slide, "logo-ai-factory", L + Inches(0.12), py + Inches(0.055), Inches(0.15))
    label(slide, L + Inches(1.68), py + Inches(0.04), Inches(7.0), [
        ("with NVIDIA — what every turn above draws on. One release manifest, one support contract.",
         6.8, True, GREY_99),
    ])
    seg(slide, Inches(4.00), py, Inches(4.00), mbot, NVIDIA_GREEN)
    label(slide, Inches(4.10), mbot + Inches(0.015), Inches(4.2), [
        ("every turn · foundational_rag over HTTP", 5.8, False, NVIDIA_GREEN)])

    gw = (CW - Inches(0.24) - 2 * Inches(0.12)) / 3
    for g, (title, gsub, rows) in enumerate(PLANES):
        gx = L + Inches(0.12) + g * (gw + Inches(0.12))
        box(slide, gx, Inches(5.74), gw, Inches(1.30), line=GREY_3E, lw=Pt(0.75))
        label(slide, gx + Inches(0.10), Inches(5.79), gw - Inches(0.20), [
            (title, 7.5, True, NVIDIA_GREEN),
            (gsub, 5.8, False, GREY_6F),
        ])
        rw = (gw - Inches(0.16)) / len(rows)
        for i, (ico, name, sub) in enumerate(rows):
            tile(slide, int(gx + Inches(0.08) + i * rw + rw / 2), Inches(6.14), int(rw),
                 ico, name, sub, NVIDIA_GREEN, GREY_99, ih=Inches(0.36),
                 name_pt=7.0, sub_pt=5.5)

    # the flywheel closes: the human verdict comes back as a weight
    fb = Inches(13.16)
    seg(slide, ecx, Inches(4.46), fb, Inches(4.46), PERSIMMON, lw=Pt(1.0), dash=True, head=False)
    seg(slide, fb, Inches(4.46), fb, Inches(6.30), PERSIMMON, lw=Pt(1.0), dash=True, head=False)
    seg(slide, fb, Inches(6.30), L + Inches(0.12) + 2 * (gw + Inches(0.12)) + gw,
        Inches(6.30), PERSIMMON, lw=Pt(1.0), dash=True)
    label(slide, Inches(9.90), mbot + Inches(0.015), Inches(3.13), [
        ("the human verdict · accepted trajectories", 5.8, False, PERSIMMON)],
        align=PP_ALIGN.RIGHT)

    cust_cx = int(L + Inches(0.12) + 2 * (gw + Inches(0.12)) + Inches(0.08)
                  + 3.5 * (gw - Inches(0.16)) / 4)
    # along the floor of the panel, below the three groups: Customizer -> MODELS
    seg(slide, cust_cx, Inches(7.30), Inches(2.20), Inches(7.30), NVIDIA_GREEN,
        lw=Pt(1.0), dash=True)
    label(slide, Inches(2.30), Inches(7.16), Inches(6.00), [
        ("LoRA on the 30B, folded back weekly — the Ultra/Lightning split is why it pays",
         5.8, False, NVIDIA_GREEN)])

    # --- infrastructure ------------------------------------------------------
    ny, nh = Inches(7.58), Inches(0.44)
    nw = (CW - 4 * Inches(0.10)) / 5
    for i, (ico, name, sub, colour) in enumerate(INFRA):
        x = L + i * (nw + Inches(0.10))
        box(slide, x, ny, nw, nh,
            fill=base.MIDNIGHT_DEEP if colour is WATERHOLE_MID else PANEL,
            line=colour, lw=Pt(0.9))
        icon(slide, ico, x + Inches(0.10), ny + Inches(0.09), Inches(0.26))
        label(slide, x + Inches(0.42), ny + Inches(0.08), nw - Inches(0.52), [
            (name, 7.0, True, colour),
            (sub, 5.5, False, GREY_99),
        ])

    # --- legend and the honest footer ----------------------------------------
    ly = Inches(8.14)
    x = L
    for colour, text in LEGEND:
        sw_ = box(slide, x, ly + Inches(0.012), Inches(0.09), Inches(0.09), fill=colour)
        sw_.line.fill.background()
        label(slide, x + Inches(0.14), ly, Inches(2.0), [(text, 5.8, False, GREY_BA)])
        x += Inches(0.14) + Inches(0.042) * len(text) + Inches(0.16)
    label(slide, Inches(5.60), ly, Inches(7.43), [
        ("Reference architecture. OpenShell, NemoClaw and the sandbox lifecycle are verified on a live "
         "SUSE cluster; the SecOps agent layer, the data flywheel and the three portfolio integrations "
         "are designed, not built.", 5.5, False, GREY_6F)], align=PP_ALIGN.RIGHT)

    prs.save(str(OUT))
    print(f"{len(slide.shapes)} shapes -> {OUT}")


if __name__ == "__main__":
    build()
